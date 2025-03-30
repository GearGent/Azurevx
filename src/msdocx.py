#!/usr/bin/python3
# BSD 3-Clause License
#
# Copyright (c) 2025, Geoffrey Argence <argence.geoffrey@gmail.com>
#
# Redistribution and use in source and binary forms, with or without
# modification, are permitted provided that the following conditions are met:
#
# 1. Redistributions of source code must retain the above copyright notice, this
#    list of conditions and the following disclaimer.
#
# 2. Redistributions in binary form must reproduce the above copyright notice,
#    this list of conditions and the following disclaimer in the documentation
#    and/or other materials provided with the distribution.
#
# 3. Neither the name of the copyright holder nor the names of its
#    contributors may be used to endorse or promote products derived from
#    this software without specific prior written permission.
#
# THIS SOFTWARE IS PROVIDED BY THE COPYRIGHT HOLDERS AND CONTRIBUTORS "AS IS"
# AND ANY EXPRESS OR IMPLIED WARRANTIES, INCLUDING, BUT NOT LIMITED TO, THE
# IMPLIED WARRANTIES OF MERCHANTABILITY AND FITNESS FOR A PARTICULAR PURPOSE ARE
# DISCLAIMED. IN NO EVENT SHALL THE COPYRIGHT HOLDER OR CONTRIBUTORS BE LIABLE
# FOR ANY DIRECT, INDIRECT, INCIDENTAL, SPECIAL, EXEMPLARY, OR CONSEQUENTIAL
# DAMAGES (INCLUDING, BUT NOT LIMITED TO, PROCUREMENT OF SUBSTITUTE GOODS OR
# SERVICES; LOSS OF USE, DATA, OR PROFITS; OR BUSINESS INTERRUPTION) HOWEVER
# CAUSED AND ON ANY THEORY OF LIABILITY, WHETHER IN CONTRACT, STRICT LIABILITY,
# OR TORT (INCLUDING NEGLIGENCE OR OTHERWISE) ARISING IN ANY WAY OUT OF THE USE
# OF THIS SOFTWARE, EVEN IF ADVISED OF THE POSSIBILITY OF SUCH DAMAGE.

import fetcher
import argparse
import docx
import locale
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches
from typing import *
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime
from lxml import etree

parser = argparse.ArgumentParser(description="Export Azure Devops reviews to a microsoft docx document")

parser.add_argument("url", type=str, help="Pull request or repository url")

parser.add_argument("-p", "--password", type=str, help="Token API to use to access the server")
parser.add_argument("-o", "--output", type=str, default="review.docx",help="Output file (.docx)")
parser.add_argument("-t", "--template", type=str, help="Template file (.docx)")

args = parser.parse_args()

if args.url:

    # Sometimes when review lines are too long, dimensions set for tables are no respected.
    # However, when manually editing the dimensions in ms Word, these remain stable after save.
    # After comparing docx of a table before and after manual editing, turned out that adding the tblLayout=fixed attribute to the parent element (element containing the table) locks the dimensions.
    # This function adds the tblLayout=fixed attribute to the tblPr of an element.
    def set_fixed_layout(table: docx.table.Table):
        tbl = table._element
        tbl_layout = etree.Element("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblLayout")
        tbl_layout.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type", "fixed")
        tbl_pr = tbl.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblPr")
        tbl_pr.append(tbl_layout)

    def print_quote(quote: fetcher.ThreadContextQuote, where: docx.document.Document | docx.table._Cell, suggestion: Optional[List[str]] = None, linesabove: int = 2, linesafter: int = 2):
        tab = where.add_table(0, 2)
        set_fixed_layout(tab)
        tab.style = "CodeTable"
        tab.autofit = False
        linenum: int = max(quote.startLine - linesabove + 1, 1)

        row = tab.add_row()
        total_width = (row.cells[0].width + row.cells[1].width)
        width_col0 = int(total_width / 100 * 10)
        width_col1 = int(total_width / 100 * 87.5)
        tab._tbl.remove(tab.rows[0]._tr)

        row = None
        for line in quote.getLinesBeforeContext(linesabove):
            row = tab.add_row()
            row.cells[0].paragraphs[-1].style = "LineNumber"
            row.cells[0].paragraphs[-1].add_run(str(linenum))
            row.cells[1].paragraphs[-1].style = "Code"
            row.cells[1].paragraphs[-1].add_run(line)
            linenum += 1

        row = tab.add_row()
        row.cells[0].paragraphs[-1].style = "LineNumber"
        row.cells[0].paragraphs[-1].add_run(str(linenum))
        context: List[str]
        if suggestion != None:
            context = suggestion
        else:
            context = quote.getContext()
        cbc = quote.getCharsBeforeContext()
        cac = quote.getCharsAfterContext()
        p = None
        if cbc or cac or len(context) > 0:
            p = row.cells[1].paragraphs[-1]
            p.style = "Code"
            p.add_run(cbc)
            if len(context) > 0:
                p.add_run(context[0]).font.highlight_color = WD_COLOR_INDEX.TURQUOISE
            linenum += 1
        for line in context[1:]:
            row = tab.add_row()
            row.cells[0].paragraphs[-1].style = "LineNumber"
            row.cells[0].paragraphs[-1].add_run(str(linenum))
            p = row.cells[1].paragraphs[-1]
            p.style = "Code"
            p.add_run(line).font.highlight_color = WD_COLOR_INDEX.TURQUOISE
            linenum += 1
        assert(p or not cac), "Script Error: A paragraph should have been created if there are characters after context on last context line"
        if cac:
            p.add_run(cac)

        for line in quote.getLinesAfterContext(linesafter):
            row = tab.add_row()
            row.cells[0].paragraphs[-1].style = "LineNumber"
            row.cells[0].paragraphs[-1].add_run(str(linenum))
            row.cells[1].paragraphs[-1].style = "Code"
            row.cells[1].paragraphs[-1].add_run(line)
            linenum += 1

        tab.columns[0].width = width_col0
        tab.columns[1].width = width_col1
        for row in tab.rows:
            row.cells[0].width = width_col0
            row.cells[1].width = width_col1

    def print_contextQuote(self, cell: docx.table._Cell, linesabove: int = 2, linesafter: int = 2):
        print_quote(quote = self, where = cell, suggestion = None, linesabove = linesabove, linesafter = linesafter)

    def print_context(self, cell: docx.table._Cell):
        p = cell.paragraphs[-1] # Table cells have a first paragraph already
        p.style = "ContextData"
        p.add_run(f"commit: {self.commit}")
        cell.add_paragraph(f"file: {self.filePath}", style = "ContextData")
        if self.quote:
            self.quote.print(cell)

    def print_comment(self, cell: docx.table._Cell, quote: fetcher.ThreadContextQuote):
        cell.add_paragraph(f"{self.author}, {datetime.fromisoformat(self.pubDate[0:len("YYYY-MM-DDThh:mm")]).strftime("%c")}:", style = "CommentMetaData")
        tab = cell.add_table(1, 1)
        set_fixed_layout(tab)
        tab.style = "Table Grid"
        comcell = tab.cell(0, 0)
        comcell.paragraphs[-1].style = "Tiny" # Table cells have a first paragraph already
        suggestion: Optional[List[str]] = None
        for line in self.content.splitlines():
            match line:
                case "```suggestion":
                    suggestion = list()
                case "```":
                    if suggestion != None:
                        comcell.add_paragraph("Suggestion:", style = "Comment")
                        print_quote(quote = quote, where = comcell, suggestion = suggestion)
                        suggestion = None
                    else:
                        comcell.add_paragraph(line, style = "Comment")
                case _:
                    if suggestion != None:
                        suggestion.append(line)
                    else:
                        comcell.add_paragraph(line, style = "Comment")

        comcell.add_paragraph().style = "Tiny"

    def print_thread(self, doc: Document):
        tab = doc.add_table(1, 1, style = "Table Grid")
        set_fixed_layout(tab)
        cell = tab.cell(0, 0)
        contextQuote = None
        if self.context:
            self.context.print(cell)
            contextQuote = self.context.quote
        for comment in self.comments:
            comment.print(cell, contextQuote)
            cell.add_paragraph(style = "Tiny")
        doc.add_paragraph()

    def print_pullrequest(self, doc: Document):
        for thread in self.get():
            thread.print(doc)

    def print_repository(self, doc: Document):
        for pr in self.get():
            pr.print(doc)

    fetcher.ThreadContextQuote.print = print_contextQuote
    fetcher.ThreadContext.print = print_context
    fetcher.Comment.print = print_comment
    fetcher.CommentThread.print = print_thread
    fetcher.PullRequest.print = print_pullrequest
    fetcher.Repository.print = print_repository

    # Set locale for comment date format
    locale.setlocale(locale.LC_TIME, "fr_FR.UTF-8")

    doc = None
    if args.template:
        doc = Document(args.template)
    else:
        doc = Document()

    # Code Style (context quote)
    style = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Consolas"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(0)

    # LineNumer Style (context quote)
    style = doc.styles.add_style("LineNumber", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Consolas"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(191, 191, 191)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style.paragraph_format.right_indent = Inches(0.05)

    # ContextData Style (context commit and file path)
    style = doc.styles.add_style("ContextData", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(166, 166, 166)

    # CommentMetaData Style (username and date of comment)
    style = doc.styles.add_style("CommentMetaData", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(68, 114, 196)

    # Comment Style
    style = doc.styles.add_style("Comment", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0, 0, 0)

    # Tiny Style (used to make small blanck separators)
    style = doc.styles.add_style("Tiny", WD_STYLE_TYPE.PARAGRAPH)
    style.font.size = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    # CodeTable Style (context quotes table style)
    style = doc.styles.add_style("CodeTable", WD_STYLE_TYPE.TABLE)
    table_style = style.element
    tbl_pr = OxmlElement("w:tblPr")
    tbl_borders = OxmlElement("w:tblBorders")
    for border_name in ["top", "left", "bottom", "right", "insideV"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        tbl_borders.append(border)
    border = OxmlElement("w:insideH")
    border.set(qn("w:val"), "nil")
    tbl_borders.append(border)
    tbl_pr.append(tbl_borders)
    table_style.append(tbl_pr)



    if "pullrequest" in args.url:
        prmd = fetcher.PullRequestMetaData.from_url(
            url = args.url,
            password = args.password
        )
        fetcher.PullRequest(prmd).print(doc)
    else:
        rmd = fetcher.RepositoryMetaData.from_url(
            url = args.url,
            password = args.password
        )
        fetcher.Repository(rmd).print(doc)

    if args.output.endswith(".docx"):
        doc.save(args.output)
    else:
        doc.save(args.output + ".docx")
